"""Search a project's own files: what's in the PDFs, decks, spreadsheets and notes the owner added.

Files are read into passages (a few paragraphs each) and ranked with BM25, the ranking search engines
were built on: no model to download, fast on any computer, and every hit points back to its file and
place so answers can cite it. The index is rebuilt only for files that changed.
"""
from __future__ import annotations

import json
import math
import re
from collections import Counter
from pathlib import Path

TEXT = {".txt", ".md", ".csv", ".json", ".html", ".htm", ".py", ".js", ".css", ".gd", ".scad", ".svg"}
WORD = re.compile(r"[a-z0-9]+(?:'[a-z]+)?")
STOP = set("a an and are as at be but by for from has have he her his i if in into is it its of on or our she so that the their them "
           "they this to was we were what when which who will with you your".split())
CHUNK = 900


def words(text):
    return [w for w in WORD.findall(text.lower()) if w not in STOP]


def extract(path):
    """The readable text of a file, as (place, text) pieces. Unknown or unreadable files give nothing."""
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext in TEXT:
            raw = path.read_text("utf-8", "replace")
            if ext in {".html", ".htm", ".svg"}:
                raw = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", raw, flags=re.S | re.I)
                raw = re.sub(r"<[^>]+>", " ", raw)
            return [("", raw)]
        if ext == ".docx":
            from docx import Document
            return [("", "\n".join(p.text for p in Document(str(path)).paragraphs))]
        if ext == ".pptx":
            from pptx import Presentation
            out = []
            for number, slide in enumerate(Presentation(str(path)).slides, 1):
                bits = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
                if slide.has_notes_slide:
                    bits.append(slide.notes_slide.notes_text_frame.text)
                out.append((f"slide {number}", "\n".join(bits)))
            return out
        if ext == ".pdf":
            from pypdf import PdfReader
            return [(f"page {i}", page.extract_text() or "") for i, page in enumerate(PdfReader(str(path)).pages, 1)]
        if ext == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(str(path), read_only=True, data_only=True)
            out = []
            for sheet in book.worksheets:
                rows = [", ".join("" if v is None else str(v) for v in row) for row in sheet.iter_rows(values_only=True, max_row=2000)]
                out.append((f"sheet {sheet.title}", "\n".join(rows)))
            return out
    except Exception:
        return []
    return []


def passages(text, size=CHUNK):
    """Paragraph-aligned pieces of about `size` characters, overlapping by one paragraph."""
    paras = [p.strip() for p in re.split(r"\n\s*\n|\n(?=[#*\-•])", text) if p.strip()]
    out, current = [], []
    for para in paras:
        while len(para) > size * 1.5:
            out.append(para[:size])
            para = para[size - 100:]
        current.append(para)
        if sum(len(p) for p in current) >= size:
            out.append("\n".join(current))
            current = current[-1:]
    if current and (not out or "\n".join(current) != out[-1]):
        out.append("\n".join(current))
    return out


class Index:
    def __init__(self, folder):
        self.folder = Path(folder)
        self.file = self.folder / ".luma" / "index.json"
        try:
            self.data = json.loads(self.file.read_text())
        except (OSError, ValueError):
            self.data = {"files": {}, "chunks": []}

    def refresh(self):
        """Index new and changed files; drop deleted ones."""
        seen = {}
        for path in sorted(self.folder.rglob("*")):
            rel = path.relative_to(self.folder).as_posix()
            if not path.is_file() or rel.startswith(".luma/") or path.stat().st_size > 30 * 2 ** 20:
                continue
            seen[rel] = path.stat().st_mtime
        changed = {rel for rel, mtime in seen.items() if self.data["files"].get(rel) != mtime}
        gone = set(self.data["files"]) - set(seen)
        if not changed and not gone:
            return self
        chunks = [c for c in self.data["chunks"] if c["file"] not in changed | gone]
        for rel in sorted(changed):
            for place, text in extract(self.folder / rel):
                for piece in passages(text):
                    terms = words(piece)
                    if terms:
                        chunks.append({"file": rel, "place": place, "text": piece[:1600], "terms": dict(Counter(terms)), "length": len(terms)})
        self.data = {"files": seen, "chunks": chunks}
        self.file.parent.mkdir(parents=True, exist_ok=True)
        self.file.write_text(json.dumps(self.data))
        return self

    def search(self, query, limit=6, k1=1.4, b=0.75):
        """Best passages for the query, each with its file and place (page, slide, sheet) to cite."""
        chunks = self.data["chunks"]
        terms = [t for t in dict.fromkeys(words(query))]
        if not chunks or not terms:
            return []
        avg = sum(c["length"] for c in chunks) / len(chunks)
        df = {t: sum(1 for c in chunks if t in c["terms"]) for t in terms}
        scored = []
        for chunk in chunks:
            score = 0.0
            for t in terms:
                tf = chunk["terms"].get(t, 0)
                if tf:
                    idf = math.log(1 + (len(chunks) - df[t] + 0.5) / (df[t] + 0.5))
                    score += idf * tf * (k1 + 1) / (tf + k1 * (1 - b + b * chunk["length"] / avg))
            if score > 0:
                scored.append((score, chunk))
        scored.sort(key=lambda pair: -pair[0])
        return [{"file": c["file"], "place": c["place"], "text": c["text"], "score": round(s, 3)} for s, c in scored[:limit]]
