"""Data: analyze the owner's spreadsheets, and build financial models that stay live in Excel.

The numbers are computed here in Python (so they can be checked), and the spreadsheet is written with
real formulas pointing at an Assumptions sheet, so changing one assumption in Excel or Numbers
recalculates the whole model.
"""
from __future__ import annotations

import csv
import io
import json
import math
import re
import statistics
from datetime import date

from luma.work.studios.base import Result, Studio, check, slug, text

NUMBER = re.compile(r"^\s*[-+(]?[$€£¥]?\s*[-+]?\d[\d,]*(?:\.\d+)?\s*%?\)?\s*$")


def to_number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else None
    s = str(value or "").strip()
    if not s or not NUMBER.match(s):
        return None
    negative = s.startswith("(") and s.endswith(")") or s.startswith("-")
    digits = re.sub(r"[^\d.]", "", s)
    try:
        n = float(digits)
    except ValueError:
        return None
    n = -n if negative else n
    return n / 100 if s.endswith("%") else n


def load_table(path):
    """Rows (dicts) and column names from a CSV, Excel or JSON file."""
    ext = path.suffix.lower()
    if ext == ".csv":
        raw = path.read_text("utf-8-sig", "replace")
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|") if raw.strip() else csv.excel
        rows = list(csv.DictReader(io.StringIO(raw), dialect=dialect))
        return rows, list(rows[0].keys()) if rows else []
    if ext == ".xlsx":
        from openpyxl import load_workbook
        sheet = load_workbook(str(path), read_only=True, data_only=True).worksheets[0]
        values = list(sheet.iter_rows(values_only=True, max_row=50000))
        if not values:
            return [], []
        header = [str(h if h is not None else f"Column {i + 1}") for i, h in enumerate(values[0])]
        return [dict(zip(header, row)) for row in values[1:] if any(v is not None for v in row)], header
    if ext == ".json":
        data = json.loads(path.read_text("utf-8"))
        rows = data if isinstance(data, list) else next((v for v in data.values() if isinstance(v, list)), [])
        rows = [r for r in rows if isinstance(r, dict)]
        return rows, list(dict.fromkeys(k for r in rows for k in r))
    raise ValueError("I can analyze CSV, Excel (.xlsx) or JSON tables.")


def profile(rows, columns):
    out = []
    for col in columns:
        values = [r.get(col) for r in rows]
        present = [v for v in values if v not in (None, "")]
        numbers = [n for n in (to_number(v) for v in present) if n is not None]
        info = {"column": col, "filled": len(present), "missing": len(values) - len(present)}
        if present and len(numbers) >= 0.8 * len(present):
            info.update(type="number", sum=round(sum(numbers), 4), mean=round(statistics.fmean(numbers), 4), min=min(numbers), max=max(numbers))
        else:
            counts = {}
            for v in present:
                counts[str(v)] = counts.get(str(v), 0) + 1
            info.update(type="text", distinct=len(counts), top=sorted(counts.items(), key=lambda kv: -kv[1])[:5])
        out.append(info)
    return out


def match_column(columns, wanted):
    wanted = str(wanted or "").strip().lower()
    if not wanted:
        return None
    return next((c for c in columns if c.lower() == wanted), None) or next((c for c in columns if wanted in c.lower() or c.lower() in wanted), None)


class AnalyzeData(Studio):
    name, agent, title, tier = "analyze_data", "data", "Data analysis", "standard"
    description = ("Analyze a table in the project. file: the CSV/XLSX/JSON file's path in the project. question: what they want to know. "
                   "group_by: the column to group by (or empty). measure: the number column to add up (or empty).")
    fields = {"file": "string", "question": "text", "group_by": "text", "measure": "text"}

    def fallback(self, ctx, task):
        tables = [f["path"] for f in ctx.files() if f["path"].lower().endswith((".csv", ".xlsx", ".json")) and not f["path"].startswith("analysis/")]
        return {"file": tables[0], "question": task, "group_by": "", "measure": ""} if tables else None

    def make(self, ctx, args):
        rel = text(args.get("file"), 300)
        rows, columns = load_table(ctx.path(rel))
        if not rows:
            raise ValueError(f"{rel} has no rows to analyze.")
        prof = profile(rows, columns)
        group, measure = match_column(columns, args.get("group_by")), match_column(columns, args.get("measure"))
        grouped = []
        if group:
            buckets = {}
            for r in rows:
                key = str(r.get(group) if r.get(group) not in (None, "") else "(blank)")
                b = buckets.setdefault(key, {"group": key, "rows": 0, "total": 0.0})
                b["rows"] += 1
                if measure:
                    b["total"] += to_number(r.get(measure)) or 0.0
            grouped = sorted(buckets.values(), key=lambda b: -(b["total"] if measure else b["rows"]))
        stem = "analysis/" + slug(rel.rsplit("/", 1)[-1].rsplit(".", 1)[0], "table")
        lines = [f"# Analysis of {rel}", "", f"{len(rows)} rows, {len(columns)} columns. Question: {text(args.get('question'), 400) or '(general overview)'}", "", "## Columns", ""]
        for p in prof:
            lines.append(f"- **{p['column']}** ({p['type']}, {p['filled']} filled, {p['missing']} missing): " +
                         (f"total {p['sum']:,.2f}, average {p['mean']:,.2f}, from {p['min']:,.2f} to {p['max']:,.2f}" if p["type"] == "number"
                          else f"{p['distinct']} distinct, most common: " + ", ".join(f"{k} ({v})" for k, v in p["top"])))
        if grouped:
            lines += ["", f"## {measure or 'Rows'} by {group}", "", f"| {group} | rows | {measure or ''} |", "|---|---:|---:|"]
            lines += [f"| {g['group']} | {g['rows']} | {g['total']:,.2f} |" if measure else f"| {g['group']} | {g['rows']} | |" for g in grouped[:30]]
        ctx.write(stem + "-analysis.md", "\n".join(lines) + "\n")
        files = [stem + "-analysis.md"]
        with io.StringIO() as buf:
            writer = csv.writer(buf)
            writer.writerow(["column", "type", "filled", "missing", "sum", "mean", "min", "max"])
            for p in prof:
                writer.writerow([p["column"], p["type"], p["filled"], p["missing"], p.get("sum", ""), p.get("mean", ""), p.get("min", ""), p.get("max", "")])
            ctx.write(stem + "-profile.csv", buf.getvalue())
        files.append(stem + "-profile.csv")
        try:
            from openpyxl import Workbook
            from openpyxl.chart import BarChart, Reference
            book = Workbook()
            sheet = book.active
            sheet.title = "Columns"
            sheet.append(["Column", "Type", "Filled", "Missing", "Total", "Average", "Min", "Max"])
            for p in prof:
                sheet.append([p["column"], p["type"], p["filled"], p["missing"], p.get("sum"), p.get("mean"), p.get("min"), p.get("max")])
            if grouped:
                g = book.create_sheet("By " + group[:24])
                g.append([group, "Rows", measure or ""])
                for row in grouped:
                    g.append([row["group"], row["rows"], row["total"] if measure else None])
                chart = BarChart()
                chart.title = f"{measure or 'Rows'} by {group}"
                chart.add_data(Reference(g, min_col=3 if measure else 2, min_row=1, max_row=min(len(grouped), 20) + 1), titles_from_data=True)
                chart.set_categories(Reference(g, min_col=1, min_row=2, max_row=min(len(grouped), 20) + 1))
                g.add_chart(chart, "E2")
            book.save(str(ctx.path(stem + "-summary.xlsx")))
            files.append(stem + "-summary.xlsx")
        except ImportError:
            pass
        total = next((p["sum"] for p in prof if p["column"] == measure), None) if measure else None
        return Result(files=files, summary=f"Analyzed {rel}: {len(rows)} rows, {len(columns)} columns" + (f", {measure} by {group}" if grouped and measure else ""),
                      data={"rows": len(rows), "columns": columns, "profile": prof, "group_by": group, "measure": measure, "grouped": grouped[:20], "total": total})

    def check(self, ctx, result):
        d = result.data
        grouped_total = sum(g["total"] for g in d.get("grouped", [])) if d.get("measure") else None
        out = [check("read the table", d.get("rows", 0) > 0, "No rows were read."), check("files written", all(ctx.exists(f) for f in result.files), "An output file is missing.")]
        if grouped_total is not None and d.get("total") is not None and len(d.get("grouped", [])) < 20:
            out.append(check("groups add up", abs(grouped_total - d["total"]) < 0.01 * max(1, abs(d["total"])), f"Groups total {grouped_total:,.2f} but the column totals {d['total']:,.2f}."))
        return out


def parse_costs(value):
    """'Salaries: 20000; Rent 3,500/month' -> [("Salaries", 20000.0), ("Rent", 3500.0)]."""
    costs = []
    for part in re.split(r"[;\n]+|,\s(?=[A-Za-z])", str(value or "")):
        m = re.match(r"\s*([A-Za-z][^:=$\d]{0,40}?)\s*[:=]?\s*\$?\s*([\d,]+(?:\.\d+)?)\s*(k)?", part, re.I)
        if m:
            amount = float(m.group(2).replace(",", "")) * (1000 if m.group(3) else 1)
            costs.append((m.group(1).strip().title(), amount))
    return costs[:12]


def project_months(a):
    """The model, month by month, in plain Python (the spreadsheet's formulas compute the same)."""
    months, customers, cash = [], a["start_customers"], -a["startup_cost"]
    for m in range(1, a["years"] * 12 + 1):
        if m > 1:
            customers = max(0.0, customers * (1 + a["growth"] - a["churn"]))
        revenue = customers * a["price"]
        cogs = revenue * a["cogs"]
        net = revenue - cogs - a["monthly_costs"]
        cash += net
        months.append({"month": m, "customers": customers, "revenue": revenue, "cogs": cogs, "gross": revenue - cogs, "costs": a["monthly_costs"], "net": net, "cash": cash})
    return months


class FinancialModel(Studio):
    name, agent, title = "financial_model", "data", "Financial model"
    description = ("Build a financial model. Use real numbers from the research where there are any; otherwise reasonable, stated assumptions. "
                   "price: what one customer pays per month. start_customers: paying customers in month one. monthly_growth_pct and monthly_churn_pct: "
                   "percent per month (8 means 8%). cogs_pct: cost of goods as a percent of revenue. monthly_costs: fixed costs like 'Salaries: 12000; Rent: 2500; Marketing: 3000'. "
                   "startup_cost: one-time cost to launch. years: 1 to 5. currency: like USD.")
    fields = {"name": "string", "currency": "text", "years": "integer", "price": "number", "start_customers": "number", "monthly_growth_pct": "number",
              "monthly_churn_pct": "number", "cogs_pct": "number", "monthly_costs": "text", "startup_cost": "number"}

    def fallback(self, ctx, task):
        return {"name": text(task, 60) or "Financial model", "currency": "USD", "years": 3, "price": 50, "start_customers": 20, "monthly_growth_pct": 10,
                "monthly_churn_pct": 3, "cogs_pct": 30, "monthly_costs": "Salaries: 8000; Rent: 2000; Marketing: 1500", "startup_cost": 25000, "_placeholder": True}

    @staticmethod
    def assumptions(args):
        num = lambda key, default, lo, hi: min(hi, max(lo, float(args.get(key) if isinstance(args.get(key), (int, float)) and not isinstance(args.get(key), bool) else default)))
        costs = parse_costs(args.get("monthly_costs")) or [("Operating costs", 5000.0)]
        return {"name": text(args.get("name"), 80) or "Financial model", "currency": (text(args.get("currency"), 8) or "USD").upper(),
                "years": int(num("years", 3, 1, 5)), "price": num("price", 50, 0, 1e9), "start_customers": num("start_customers", 10, 0, 1e9),
                "growth": num("monthly_growth_pct", 5, -50, 200) / 100, "churn": num("monthly_churn_pct", 2, 0, 100) / 100, "cogs": num("cogs_pct", 30, 0, 100) / 100,
                "cost_lines": costs, "monthly_costs": sum(c for _, c in costs), "startup_cost": num("startup_cost", 0, 0, 1e12)}

    def make(self, ctx, args):
        a = self.assumptions(args)
        months = project_months(a)
        years = []
        for y in range(a["years"]):
            span = months[y * 12:(y + 1) * 12]
            years.append({"year": y + 1, "revenue": sum(m["revenue"] for m in span), "gross": sum(m["gross"] for m in span), "costs": sum(m["costs"] for m in span),
                          "net": sum(m["net"] for m in span), "customers": span[-1]["customers"], "cash": span[-1]["cash"]})
        breakeven = next((m["month"] for m in months if m["net"] > 0), None)
        payback = next((m["month"] for m in months if m["cash"] >= 0), None)
        stem = "finance/" + slug(a["name"], "model")[:40]
        self.workbook(ctx.path(stem + ".xlsx"), a, months)
        with io.StringIO() as buf:
            writer = csv.writer(buf)
            writer.writerow(["Year", "Revenue", "Gross profit", "Fixed costs", "Net", "Customers at year end", "Cash at year end"])
            for y in years:
                writer.writerow([y["year"], round(y["revenue"], 2), round(y["gross"], 2), round(y["costs"], 2), round(y["net"], 2), round(y["customers"]), round(y["cash"], 2)])
            ctx.write(stem + "-yearly.csv", buf.getvalue())
        money = lambda v: f"{a['currency']} {v:,.0f}"
        md = [f"# {a['name']}", "", f"_Built {date.today():%B %d, %Y}. Open the spreadsheet and change any number on the Assumptions sheet; everything recalculates._", "",
              "## Assumptions", "", f"- Price per customer per month: {money(a['price'])}", f"- Customers in month one: {a['start_customers']:,.0f}",
              f"- Growth {a['growth'] * 100:.1f}% and churn {a['churn'] * 100:.1f}% per month", f"- Cost of goods: {a['cogs'] * 100:.0f}% of revenue",
              f"- Fixed costs per month: " + ", ".join(f"{n} {money(c)}" for n, c in a["cost_lines"]), f"- Startup cost: {money(a['startup_cost'])}", "",
              "## Results", "", "| Year | Revenue | Net | Customers | Cash |", "|---:|---:|---:|---:|---:|"]
        md += [f"| {y['year']} | {money(y['revenue'])} | {money(y['net'])} | {y['customers']:,.0f} | {money(y['cash'])} |" for y in years]
        md += ["", f"Profitable from month {breakeven}." if breakeven else "Not profitable within the model's horizon.",
               f"The startup cost is paid back in month {payback}." if payback else "The startup cost isn't paid back within the horizon."]
        if args.get("_placeholder"):
            md += ["", "**These are placeholder assumptions.** Replace them with real ones on the Assumptions sheet."]
        ctx.write(stem + ".md", "\n".join(md) + "\n")
        return Result(files=[stem + ".xlsx", stem + "-yearly.csv", stem + ".md"],
                      summary=f"Financial model: year 1 revenue {money(years[0]['revenue'])}, year {a['years']} revenue {money(years[-1]['revenue'])}; " +
                              (f"profitable from month {breakeven}." if breakeven else "not profitable within the horizon."),
                      data={"currency": a["currency"], "years": [{k: round(v, 2) if isinstance(v, float) else v for k, v in y.items()} for y in years],
                            "breakeven_month": breakeven, "payback_month": payback, "placeholder": bool(args.get("_placeholder")),
                            "assumptions": {k: v for k, v in a.items() if k != "cost_lines"} | {"cost_lines": [list(c) for c in a["cost_lines"]]}})

    @staticmethod
    def workbook(path, a, months):
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, LineChart, Reference
        from openpyxl.styles import Font, PatternFill
        book = Workbook()
        s = book.active
        s.title = "Assumptions"
        bold, fill = Font(bold=True), PatternFill("solid", fgColor="EDE9FE")
        rows = [("Price per customer per month", a["price"], "#,##0.00"), ("Customers in month one", a["start_customers"], "#,##0"),
                ("Monthly growth", a["growth"], "0.0%"), ("Monthly churn", a["churn"], "0.0%"), ("Cost of goods (share of revenue)", a["cogs"], "0.0%"),
                ("Fixed costs per month", None, "#,##0"), ("Startup cost", a["startup_cost"], "#,##0")]
        s.append(["Assumption", "Value"])
        for label, value, fmt in rows:
            s.append([label, value])
            s.cell(row=s.max_row, column=2).number_format = fmt
            s.cell(row=s.max_row, column=2).fill = fill
        s.append([])
        s.append(["Fixed cost lines", "Per month"])
        first = s.max_row + 1
        for name, amount in a["cost_lines"]:
            s.append([name, amount])
            s.cell(row=s.max_row, column=2).number_format = "#,##0"
            s.cell(row=s.max_row, column=2).fill = fill
        s["B7"] = f"=SUM(B{first}:B{s.max_row})"
        for cell in ("A1", "B1", f"A{first - 1}", f"B{first - 1}"):
            s[cell].font = bold
        s.column_dimensions["A"].width = 36
        s.column_dimensions["B"].width = 16
        m = book.create_sheet("Monthly")
        m.append(["Month", "Customers", "Revenue", "Cost of goods", "Gross profit", "Fixed costs", "Net", "Cash"])
        for i in range(len(months)):
            r = i + 2
            m.append([i + 1, "=Assumptions!B3" if i == 0 else f"=MAX(0,B{r - 1}*(1+Assumptions!$B$4-Assumptions!$B$5))", f"=B{r}*Assumptions!$B$2",
                      f"=C{r}*Assumptions!$B$6", f"=C{r}-D{r}", "=Assumptions!$B$7", f"=E{r}-F{r}", f"=G{r}-Assumptions!$B$8" if i == 0 else f"=H{r - 1}+G{r}"])
            for col in "BCDEFGH":
                m[f"{col}{r}"].number_format = "#,##0"
        y = book.create_sheet("Yearly")
        y.append(["Year", "Revenue", "Gross profit", "Fixed costs", "Net", "Customers at year end", "Cash at year end"])
        for yr in range(a["years"]):
            lo, hi = yr * 12 + 2, yr * 12 + 13
            y.append([f"Year {yr + 1}", f"=SUM(Monthly!C{lo}:C{hi})", f"=SUM(Monthly!E{lo}:E{hi})", f"=SUM(Monthly!F{lo}:F{hi})", f"=SUM(Monthly!G{lo}:G{hi})", f"=Monthly!B{hi}", f"=Monthly!H{hi}"])
            for col in "BCDEFG":
                y[f"{col}{yr + 2}"].number_format = "#,##0"
        for sheet in (m, y):
            for cell in sheet[1]:
                cell.font = bold
            for col in "ABCDEFGH":
                sheet.column_dimensions[col].width = 18
        bars = BarChart()
        bars.title = "Revenue and net by year"
        bars.add_data(Reference(y, min_col=2, min_row=1, max_row=a["years"] + 1), titles_from_data=True)
        bars.add_data(Reference(y, min_col=5, min_row=1, max_row=a["years"] + 1), titles_from_data=True)
        bars.set_categories(Reference(y, min_col=1, min_row=2, max_row=a["years"] + 1))
        y.add_chart(bars, "I2")
        line = LineChart()
        line.title = "Cash over time"
        line.add_data(Reference(m, min_col=8, min_row=1, max_row=len(months) + 1), titles_from_data=True)
        m.add_chart(line, "J2")
        path.parent.mkdir(parents=True, exist_ok=True)
        book.save(str(path))

    def check(self, ctx, result):
        d = result.data
        out = [check("assumptions are real", not d.get("placeholder"), "These are placeholder numbers; replace them on the Assumptions sheet.", blocking=False)]
        try:
            from openpyxl import load_workbook
            book = load_workbook(str(ctx.path(result.files[0])))
            monthly = book["Monthly"]
            formulas = sum(1 for row in monthly.iter_rows(min_row=2, max_row=4) for c in row if isinstance(c.value, str) and c.value.startswith("="))
            out.append(check("spreadsheet opens with live formulas", formulas >= 12 and "Yearly" in book.sheetnames, "The spreadsheet is missing its formulas."))
        except Exception as error:
            out.append(check("spreadsheet opens with live formulas", False, str(error)))
        years = d.get("years", [])
        out.append(check("numbers make sense", years and all(math.isfinite(y["revenue"]) and y["customers"] >= 0 for y in years), "The model produced impossible numbers."))
        return out
