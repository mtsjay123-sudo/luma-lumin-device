"""Documents and decks: finished Word, PDF and Markdown files, and PowerPoint decks with real charts.

Both reuse Luma's tested builders (the brain writes the content; the layout is code), and both are
checked by opening the file again the way Word, PowerPoint or a PDF reader would.
"""
from __future__ import annotations

import re

from luma.work.studios.base import Result, Studio, check, slug, text


class Document(Studio):
    name, agent, title = "document", "document", "Document"
    description = ("Write a finished document: a report, business plan, proposal, memo, guide or one-pager. format is docx, pdf or md. "
                   "Each section has a heading, full-sentence paragraphs and optional bullets. Use the real findings and numbers from earlier steps, "
                   "and cite sources like [3] where research supplied them. No placeholders.")
    fields = {"title": "string", "subtitle": "text", "format": "text", "sections": "sections"}

    def make(self, ctx, args):
        from luma.skills import documents
        fmt = str(args.get("format") or "docx").lower().strip(". ")
        fmt = fmt if fmt in {"docx", "pdf", "md"} else "docx"
        spec = documents.validate_doc({**args, "format": "pdf" if fmt == "pdf" else "docx"})
        rel = f"docs/{slug(spec['title'], 'document')}.{fmt}"
        if fmt == "md":
            lines = [f"# {spec['title']}"] + ([f"_{spec['subtitle']}_"] if spec["subtitle"] else [])
            for s in spec["sections"]:
                lines += ["", f"## {s['heading']}"] if s["heading"] else [""]
                lines += [p + "\n" for p in s["paragraphs"]] + [f"- {b}" for b in s["bullets"]]
            ctx.write(rel, "\n".join(lines) + "\n")
        else:
            (documents.build_pdf if fmt == "pdf" else documents.build_docx)(spec, ctx.path(rel))
        words = sum(len(p.split()) for s in spec["sections"] for p in s["paragraphs"] + s["bullets"])
        return Result(files=[rel], summary=f"{spec['title']}: {len(spec['sections'])} sections, about {words} words.",
                      data={"title": spec["title"], "format": fmt, "sections": [s["heading"] for s in spec["sections"]], "words": words})

    def check(self, ctx, result):
        rel, d = result.files[0], result.data
        out = [check("has real content", d.get("words", 0) >= 60, f"Only {d.get('words', 0)} words."),
               check("no placeholders", not re.search(r"\[(?:insert|your|todo|tbd|placeholder)[^\]]*\]|lorem ipsum", ctx.read_text(rel) if rel.endswith(".md") else "", re.I), "There's placeholder text left.")]
        try:
            if rel.endswith(".docx"):
                from docx import Document as Docx
                heads = [p for p in Docx(str(ctx.path(rel))).paragraphs if p.style.name.startswith("Heading")]
                out.append(check("opens in Word", heads, "The Word file has no headings."))
            elif rel.endswith(".pdf"):
                raw = ctx.path(rel).read_bytes()
                pages = None
                try:
                    from pypdf import PdfReader
                    pages = len(PdfReader(str(ctx.path(rel))).pages)
                except ImportError:
                    pages = raw.count(b"/Type /Page") - raw.count(b"/Type /Pages")
                out.append(check("opens as a PDF", raw.startswith(b"%PDF") and pages and pages > 0, "The PDF has no pages."))
        except Exception as error:
            out.append(check("opens", False, str(error)))
        return out


class Deck(Studio):
    name, agent, title = "deck", "document", "PowerPoint deck"
    description = ("Build a polished, editable PowerPoint deck. Slide titles are claims; 3 to 5 short bullets; speaker notes say what to say. "
                   "layout is title, section, bullets, big_number, quote, closing, chart, diagram or table. Use chart (chart_type bar, column, line or pie, "
                   "categories, series of numbers) for numbers from the financial model or research, diagram (steps) for how something works, "
                   "table (rows, first row is the header) for comparisons like competitors. theme is midnight, daylight, ember or forest. "
                   "Make as many slides as asked for.")
    fields = {"title": "string", "subtitle": "text", "theme": "text", "slides": "slides"}

    def make(self, ctx, args):
        from luma.skills import decks
        spec = dict(args)
        if spec.get("theme") not in decks.THEMES:
            spec["theme"] = "midnight"
        clean = decks.validate_outline(spec)
        rel = f"deck/{slug(clean['title'], 'deck')}.pptx"
        count = decks.build_deck(clean, ctx.path(rel))
        layouts = [s["layout"] for s in clean["slides"]]
        return Result(files=[rel], summary=f"{clean['title']}: a {count}-slide deck" + (f" with {layouts.count('chart')} chart{'s' if layouts.count('chart') != 1 else ''}" if "chart" in layouts else "") + ".",
                      data={"title": clean["title"], "slides": count, "layouts": layouts, "theme": clean["theme"]})

    def check(self, ctx, result):
        d = result.data
        wanted = re.search(r"\b(\d{1,2})[- ](?:slide|page)", ctx.request + " " + ctx.task, re.I)
        out = []
        try:
            from pptx import Presentation
            slides = list(Presentation(str(ctx.path(result.files[0]))).slides)
            out.append(check("opens in PowerPoint", len(slides) == d["slides"], "The saved deck doesn't match what was built."))
        except Exception as error:
            out.append(check("opens in PowerPoint", False, str(error)))
        if wanted:
            n = int(wanted.group(1))
            out.append(check(f"{n} slides as asked", abs(d["slides"] - n) <= max(1, n // 10), f"Asked for {n} slides, built {d['slides']}."))
        out.append(check("uses the numbers", "chart" in d.get("layouts", []) or not ctx.has_input("financial_model"), "There's a financial model but no chart of it.", blocking=False))
        return out
