"""Build real, good-looking .pptx decks from a structured outline.

The model writes the content (titles, bullets, notes); this module owns the
layout so every deck comes out clean and consistent, whoever wrote it.
"""
from __future__ import annotations

import re
from pathlib import Path

THEMES = {
    "midnight": {"bg": "0F1020", "panel": "1B1D36", "text": "F4F2FF", "muted": "A6A3C8", "accent": "8C7BFF", "accent2": "4FD1C5"},
    "daylight": {"bg": "FBFAF7", "panel": "F0EEE8", "text": "17171C", "muted": "6B6A72", "accent": "5B4BDB", "accent2": "E0754F"},
    "ember": {"bg": "1A1210", "panel": "2A1D19", "text": "FFF4EE", "muted": "C7A99C", "accent": "FF7A45", "accent2": "FFC15E"},
    "forest": {"bg": "0E1A16", "panel": "16281F", "text": "EEF7F1", "muted": "9DB8A8", "accent": "4CC38A", "accent2": "D7E36B"},
}
LAYOUTS = ("title", "section", "bullets", "big_number", "two_column", "quote", "closing")
HEAD_FONT = "Helvetica Neue"
BODY_FONT = "Helvetica Neue"


def _clean(value, label, limit, required=True):
    if value is None and not required:
        return ""
    if not isinstance(value, str) or (required and not value.strip()):
        raise ValueError(f"{label} needs text.")
    value = " ".join(value.split())
    if len(value) > limit:
        raise ValueError(f"{label} is too long for a slide (keep it under {limit} characters).")
    return value


def _bullets(value, label, required=False):
    if value is None:
        if required: raise ValueError(f"{label} needs at least one bullet.")
        return []
    if not isinstance(value, list) or not all(isinstance(v, str) for v in value):
        raise ValueError(f"{label} must be a list of short lines.")
    items = [_clean(v, label, 180) for v in value if v.strip()]
    if required and not items:
        raise ValueError(f"{label} needs at least one bullet.")
    if len(items) > 6:
        raise ValueError(f"{label} has more than six bullets; split it into two slides.")
    return items


def validate_outline(spec):
    """Check the whole outline before any file is written."""
    if not isinstance(spec, dict):
        raise ValueError("A deck needs a title and slides.")
    title = _clean(spec.get("title"), "Deck title", 90)
    subtitle = _clean(spec.get("subtitle"), "Deck subtitle", 160, required=False)
    theme = spec.get("theme") or "midnight"
    if theme not in THEMES:
        raise ValueError("Choose a theme: " + ", ".join(THEMES) + ".")
    slides = spec.get("slides")
    if not isinstance(slides, list) or not 1 <= len(slides) <= 30:
        raise ValueError("A deck needs between 1 and 30 slides.")
    clean = []
    for index, slide in enumerate(slides, 1):
        where = f"Slide {index}"
        if not isinstance(slide, dict):
            raise ValueError(f"{where} must be an object.")
        layout = slide.get("layout") or "bullets"
        if layout not in LAYOUTS:
            raise ValueError(f"{where} has an unknown layout; use one of {', '.join(LAYOUTS)}.")
        row = {"layout": layout,
               "title": _clean(slide.get("title"), f"{where} title", 110, required=layout not in {"quote"}),
               "subtitle": _clean(slide.get("subtitle"), f"{where} subtitle", 200, required=False),
               "notes": _clean(slide.get("notes"), f"{where} notes", 2500, required=False)}
        if layout == "bullets":
            row["bullets"] = _bullets(slide.get("bullets"), f"{where} bullets", required=True)
        elif layout == "big_number":
            row["number"] = _clean(slide.get("number"), f"{where} number", 16)
            row["caption"] = _clean(slide.get("caption"), f"{where} caption", 200, required=False)
            row["bullets"] = _bullets(slide.get("bullets"), f"{where} bullets")[:3]
        elif layout == "two_column":
            row["left_title"] = _clean(slide.get("left_title"), f"{where} left heading", 60)
            row["right_title"] = _clean(slide.get("right_title"), f"{where} right heading", 60)
            row["left_bullets"] = _bullets(slide.get("left_bullets"), f"{where} left bullets", required=True)
            row["right_bullets"] = _bullets(slide.get("right_bullets"), f"{where} right bullets", required=True)
        elif layout == "quote":
            row["quote"] = _clean(slide.get("quote"), f"{where} quote", 280)
            row["attribution"] = _clean(slide.get("attribution"), f"{where} attribution", 100, required=False)
        elif layout in {"closing", "section", "title"}:
            row["bullets"] = _bullets(slide.get("bullets"), f"{where} bullets")[:4]
        clean.append(row)
    return {"title": title, "subtitle": subtitle, "theme": theme, "slides": clean}


def safe_filename(title):
    stem = re.sub(r"[^A-Za-z0-9]+", "-", title).strip("-")[:60] or "deck"
    return stem + ".pptx"


def build_deck(spec, path):
    """Render a validated outline to `path`. Returns slide count."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt

    spec = validate_outline(spec)
    colors = THEMES[spec["theme"]]
    rgb = lambda key: RGBColor.from_string(colors[key])
    deck = Presentation()
    deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    W, H = deck.slide_width, deck.slide_height
    blank = deck.slide_layouts[6]
    inch = 914400

    def background(slide):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb("bg")

    def box(slide, x, y, w, h, text, size, color="text", bold=False, font=BODY_FONT, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, spacing=1.1):
        shape = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        para = frame.paragraphs[0]
        para.alignment = align
        para.line_spacing = spacing
        run = para.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font
        run.font.color.rgb = rgb(color)
        return frame

    def rect(slide, x, y, w, h, color, rounded=False):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, int(x), int(y), int(w), int(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        shape.line.fill.background()
        shape.shadow.inherit = False
        if rounded:
            shape.adjustments[0] = 0.08
        return shape

    def bullet_list(slide, x, y, w, h, items, size=20):
        frame = slide.shapes.add_textbox(int(x), int(y), int(w), int(h)).text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        for index, item in enumerate(items):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            para.space_after = Pt(size * 0.75)
            para.line_spacing = 1.12
            dot = para.add_run()
            dot.text = "●  "
            dot.font.size = Pt(size * 0.55)
            dot.font.color.rgb = rgb("accent")
            dot.font.name = BODY_FONT
            run = para.add_run()
            run.text = item
            run.font.size = Pt(size)
            run.font.name = BODY_FONT
            run.font.color.rgb = rgb("text")
        return frame

    def heading(slide, title, kicker=""):
        """Kicker label above a claim-style title; returns where content can start."""
        rect(slide, 0.7 * inch, 0.55 * inch, 0.5 * inch, 0.07 * inch, "accent")
        top = 0.8 * inch
        if kicker:
            box(slide, 0.7 * inch, top, W - 1.4 * inch, 0.35 * inch, kicker.upper(), 12, color="accent2", bold=True)
            top += 0.4 * inch
        lines = 1 if len(title) <= 40 else 2 if len(title) <= 80 else 3
        box(slide, 0.7 * inch, top, W - 1.4 * inch, lines * 0.62 * inch, title, 34 if lines < 3 else 28, bold=True, font=HEAD_FONT, spacing=1.0)
        return top + lines * 0.62 * inch + 0.45 * inch

    def footer(slide, number):
        box(slide, 0.7 * inch, H - 0.55 * inch, 6 * inch, 0.3 * inch, spec["title"], 10, color="muted")
        box(slide, W - 1.7 * inch, H - 0.55 * inch, 1.0 * inch, 0.3 * inch, str(number), 10, color="muted", align=PP_ALIGN.RIGHT)

    sections = 0
    for number, s in enumerate(spec["slides"], 1):
        slide = deck.slides.add_slide(blank)
        background(slide)
        layout = s["layout"]
        if layout == "title":
            rect(slide, 0, 0, 0.18 * inch, H, "accent")
            box(slide, 1.0 * inch, 2.2 * inch, W - 2.0 * inch, 1.6 * inch, s["title"], 54, bold=True, font=HEAD_FONT, anchor=MSO_ANCHOR.BOTTOM, spacing=0.95)
            sub = s["subtitle"] or spec["subtitle"]
            if sub:
                box(slide, 1.0 * inch, 4.0 * inch, W - 2.5 * inch, 1.0 * inch, sub, 22, color="muted")
            if s.get("bullets"):
                box(slide, 1.0 * inch, H - 1.2 * inch, W - 2.0 * inch, 0.4 * inch, "  ·  ".join(s["bullets"]), 13, color="accent2")
        elif layout == "section":
            sections += 1
            box(slide, 1.0 * inch, 1.2 * inch, 3 * inch, 1.6 * inch, f"{sections:02d}", 96, color="accent", bold=True, font=HEAD_FONT)
            box(slide, 1.0 * inch, 3.1 * inch, W - 2.0 * inch, 1.4 * inch, s["title"], 44, bold=True, font=HEAD_FONT)
            if s["subtitle"]:
                box(slide, 1.0 * inch, 4.5 * inch, W - 2.0 * inch, 0.8 * inch, s["subtitle"], 20, color="muted")
        elif layout == "bullets":
            top = heading(slide, s["title"], s["subtitle"])
            size = 22 if len(s["bullets"]) <= 4 else 19
            bullet_list(slide, 0.75 * inch, top, W - 1.6 * inch, H - top - 0.9 * inch, s["bullets"], size)
            footer(slide, number)
        elif layout == "big_number":
            top = max(heading(slide, s["title"], s["subtitle"]), 2.3 * inch)
            box(slide, 0.7 * inch, top, 6.2 * inch, 2.0 * inch, s["number"], 110, color="accent", bold=True, font=HEAD_FONT, anchor=MSO_ANCHOR.MIDDLE, spacing=0.9)
            if s["caption"]:
                box(slide, 0.75 * inch, top + 2.1 * inch, 6.0 * inch, 1.0 * inch, s["caption"], 20, color="muted")
            if s.get("bullets"):
                rect(slide, 7.3 * inch, top + 0.05 * inch, W - 8.0 * inch, 0.75 * inch + 0.62 * inch * len(s["bullets"]), "panel", rounded=True)
                bullet_list(slide, 7.65 * inch, top + 0.4 * inch, W - 8.7 * inch, 0.62 * inch * len(s["bullets"]), s["bullets"], 17)
            footer(slide, number)
        elif layout == "two_column":
            top = heading(slide, s["title"], s["subtitle"])
            col = (W - 1.4 * inch - 0.4 * inch) / 2
            rows = max(len(s["left_bullets"]), len(s["right_bullets"]))
            height = min(H - top - 0.85 * inch, 1.35 * inch + 0.55 * inch * rows)
            for i, (head, items, color) in enumerate(((s["left_title"], s["left_bullets"], "accent"), (s["right_title"], s["right_bullets"], "accent2"))):
                x = 0.7 * inch + i * (col + 0.4 * inch)
                rect(slide, x, top, col, height, "panel", rounded=True)
                box(slide, x + 0.35 * inch, top + 0.3 * inch, col - 0.7 * inch, 0.5 * inch, head, 20, color=color, bold=True, font=HEAD_FONT)
                bullet_list(slide, x + 0.35 * inch, top + 0.95 * inch, col - 0.7 * inch, H - top - 2.0 * inch, items, 17)
            footer(slide, number)
        elif layout == "quote":
            box(slide, 1.0 * inch, 0.95 * inch, 2 * inch, 1.2 * inch, "“", 120, color="accent", bold=True, font=HEAD_FONT)
            box(slide, 1.4 * inch, 2.45 * inch, W - 2.8 * inch, 2.4 * inch, s["quote"], 30, font=HEAD_FONT, spacing=1.15)
            if s["attribution"]:
                box(slide, 1.4 * inch, 5.0 * inch, W - 2.8 * inch, 0.5 * inch, "— " + s["attribution"], 18, color="muted")
            if s["title"]:
                box(slide, 1.4 * inch, 0.75 * inch, W - 2.8 * inch, 0.4 * inch, s["title"].upper(), 12, color="muted", bold=True)
            footer(slide, number)
        elif layout == "closing":
            rect(slide, 0, H - 0.18 * inch, W, 0.18 * inch, "accent")
            box(slide, 1.0 * inch, 1.8 * inch, W - 2.0 * inch, 1.6 * inch, s["title"], 52, bold=True, font=HEAD_FONT, anchor=MSO_ANCHOR.BOTTOM)
            if s["subtitle"]:
                box(slide, 1.0 * inch, 3.6 * inch, W - 2.0 * inch, 0.9 * inch, s["subtitle"], 22, color="muted")
            if s.get("bullets"):
                box(slide, 1.0 * inch, 4.6 * inch, W - 2.0 * inch, 0.9 * inch, "   ·   ".join(s["bullets"]), 16, color="accent2")
        if s["notes"]:
            slide.notes_slide.notes_text_frame.text = s["notes"]
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(".partial.pptx")
    deck.save(str(temporary))
    temporary.replace(target)
    return len(spec["slides"])
