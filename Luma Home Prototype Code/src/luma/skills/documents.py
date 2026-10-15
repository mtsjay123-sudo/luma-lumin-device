"""Luma makes things: Word documents, PDFs, and small apps you can use right away.

A document is a title and sections (a heading, paragraphs, bullets) that Luma writes;
it becomes a .docx (opens in Word or Pages) or a PDF. A small app is one web page
built from a tested template (a tracker, a checklist, flashcards, a countdown) with
Luma's content dropped in as data, so it works the first time, offline, forever.
"""
from __future__ import annotations

import html
import json
import os
import re
from datetime import datetime

KINDS = ("tracker", "checklist", "flashcards", "countdown")


def _text(value, label, limit, required=True):
    value = " ".join(str(value or "").split())
    if required and not value:
        raise ValueError(f"The {label} is empty.")
    return value[:limit]


def safe_filename(title, ext):
    stem = re.sub(r"[^A-Za-z0-9]+", "-", title).strip("-")[:60] or "luma"
    return f"{stem}.{ext}"


# -- documents -----------------------------------------------------------------
def validate_doc(spec):
    title = _text(spec.get("title"), "title", 140)
    fmt = str(spec.get("format") or "docx").lower().strip(". ")
    fmt = "pdf" if fmt == "pdf" else "docx"
    sections = []
    for raw in (spec.get("sections") or [])[:30]:
        if not isinstance(raw, dict):
            continue
        heading = _text(raw.get("heading"), "heading", 140, required=False)
        paragraphs = [_text(p, "paragraph", 2400, required=False) for p in (raw.get("paragraphs") or [])[:12]]
        bullets = [_text(b, "bullet", 400, required=False) for b in (raw.get("bullets") or [])[:16]]
        paragraphs, bullets = [p for p in paragraphs if p], [b for b in bullets if b]
        if heading or paragraphs or bullets:
            sections.append({"heading": heading, "paragraphs": paragraphs, "bullets": bullets})
    if not sections:
        raise ValueError("The document needs at least one section with something in it.")
    return {"title": title, "subtitle": _text(spec.get("subtitle"), "subtitle", 200, required=False), "format": fmt, "sections": sections}


def build_docx(spec, path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Helvetica Neue"
    style.font.size = Pt(11)
    head = doc.add_heading(spec["title"], level=0)
    for run in head.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x1B, 0x2E)
    if spec.get("subtitle"):
        sub = doc.add_paragraph(spec["subtitle"])
        sub.runs[0].italic = True
    for section in spec["sections"]:
        if section["heading"]:
            doc.add_heading(section["heading"], level=1)
        for paragraph in section["paragraphs"]:
            doc.add_paragraph(paragraph)
        for bullet in section["bullets"]:
            doc.add_paragraph(bullet, style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def doc_html(spec):
    """The document as a clean, printable page (the PDF is printed from this)."""
    parts = [f"<h1>{html.escape(spec['title'])}</h1>"]
    if spec.get("subtitle"):
        parts.append(f"<p class=sub>{html.escape(spec['subtitle'])}</p>")
    for s in spec["sections"]:
        if s["heading"]:
            parts.append(f"<h2>{html.escape(s['heading'])}</h2>")
        parts += [f"<p>{html.escape(p)}</p>" for p in s["paragraphs"]]
        if s["bullets"]:
            parts.append("<ul>" + "".join(f"<li>{html.escape(b)}</li>" for b in s["bullets"]) + "</ul>")
    return ("<!doctype html><meta charset=utf-8><title>" + html.escape(spec["title"]) + "</title><style>"
            "@page{margin:22mm 20mm}body{font:11.5pt/1.55 -apple-system,'Helvetica Neue',Helvetica,Arial,sans-serif;color:#1f1b2e}"
            "h1{font-size:24pt;margin:0 0 4pt;letter-spacing:-.01em}h2{font-size:14pt;margin:18pt 0 6pt;color:#3b2f7a}"
            ".sub{color:#6b6680;margin:0 0 14pt}ul{padding-left:18pt}li{margin:3pt 0}p{margin:0 0 8pt}"
            "</style>" + "".join(parts))


def build_pdf(spec, path):
    """A good-looking PDF printed by a browser; a plain one written directly if no browser is around."""
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            browser = p.chromium.launch(executable_path=os.environ.get("LUMA_BROWSER_EXECUTABLE") or None)
            try:
                page = browser.new_page()
                page.set_content(doc_html(spec), wait_until="load")
                page.pdf(path=str(path), format="Letter", print_background=True)
            finally:
                browser.close()
        if path.exists() and path.stat().st_size > 500:
            return path
    except Exception:
        pass
    path.write_bytes(simple_pdf(spec))
    return path


def _width(ch):
    if ch in "il.,;:'|!jt ()[]": return 0.30
    if ch in "mwMW@": return 0.85
    if ch.isupper(): return 0.68
    if ch.isdigit(): return 0.56
    return 0.52


def _wrap(text, size, room):
    lines, line, used = [], "", 0.0
    for word in text.split():
        w = sum(_width(c) for c in word + " ") * size
        if line and used + w > room:
            lines.append(line.rstrip()); line, used = "", 0.0
        line += word + " "; used += w
    if line.strip():
        lines.append(line.rstrip())
    return lines


def simple_pdf(spec):
    """A dependency-free PDF: Helvetica, headings, paragraphs, bullets, page breaks."""
    def esc(s):
        s = s.encode("latin-1", "replace").decode("latin-1")
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    W, H, M = 612, 792, 64
    pages, ops, y = [], [], H - M

    def put(text, font, size, gap_before=0, indent=0, bullet=False):
        nonlocal ops, y
        y -= gap_before
        for i, line in enumerate(_wrap(text, size, W - 2 * M - indent) or [""]):
            if y < M + size:
                pages.append(ops); ops = []; y = H - M
            y -= size * 1.35
            prefix = "- " if bullet and i == 0 else ("  " if bullet else "")
            ops.append(f"BT /{font} {size} Tf {M + indent} {y:.1f} Td ({esc(prefix + line)}) Tj ET")

    put(spec["title"], "F2", 22)
    if spec.get("subtitle"): put(spec["subtitle"], "F1", 12, 4)
    for s in spec["sections"]:
        if s["heading"]: put(s["heading"], "F2", 14, 14)
        for p in s["paragraphs"]: put(p, "F1", 11, 5)
        for b in s["bullets"]: put(b, "F1", 11, 2, indent=10, bullet=True)
    pages.append(ops)
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
               "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>"]
    kids = []
    for page_ops in pages:
        stream = "\n".join(page_ops).encode("latin-1", "replace")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1") + stream + b"\nendstream")
        content = len(objects)
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {W} {H}] /Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> /Contents {content} 0 R >>")
        kids.append(len(objects))
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>"
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        body = obj if isinstance(obj, bytes) else obj.encode("latin-1")
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


# -- small apps -------------------------------------------------------------------
def _list(value, limit, item_limit=120):
    if isinstance(value, list):
        items = value
    else:
        items = re.split(r"\s*(?:\n|;|,(?![^()]*\)))\s*", str(value or ""))
    return [_text(i, "item", item_limit, required=False) for i in items if str(i).strip()][:limit]


def validate_app(spec):
    title = _text(spec.get("title"), "title", 80)
    kind = str(spec.get("kind") or "").lower().strip()
    if kind not in KINDS:
        kind = "tracker"
    out = {"title": title, "kind": kind}
    if kind == "tracker":
        columns = _list(spec.get("columns"), 8, 40) or ["Date", "What", "Amount"]
        out["columns"] = [{"name": c, "type": "number" if re.search(r"amount|price|cost|\$|total|hours?|minutes?|miles?|km|reps|weight|lbs|kg|calories|count|qty|quantity", c, re.I)
                           else "date" if re.search(r"\bdate\b|\bday\b|when", c, re.I) else "text"} for c in columns]
    elif kind == "checklist":
        out["items"] = _list(spec.get("items"), 60) or ["First thing"]
    elif kind == "flashcards":
        cards = []
        for c in (spec.get("cards") or [])[:80]:
            if isinstance(c, dict) and str(c.get("front") or "").strip():
                cards.append({"front": _text(c.get("front"), "card", 300), "back": _text(c.get("back"), "answer", 600, required=False)})
        if not cards:
            raise ValueError("Flashcards need at least one card.")
        out["cards"] = cards
    else:
        when = str(spec.get("date") or "").strip()
        try:
            out["date"] = datetime.fromisoformat(when.replace("Z", "+00:00")).isoformat()
        except ValueError as error:
            raise ValueError("The countdown needs a date, like 2026-12-25.") from error
    return out


APP_CSS = """
:root{--bg:#f6f4fb;--card:rgba(255,255,255,.78);--ink:#1f1b2e;--muted:#6b6680;--line:rgba(31,27,46,.1);--a:#7c5cff;--b:#46c8ff}
@media (prefers-color-scheme:dark){:root{--bg:#121019;--card:rgba(32,28,44,.8);--ink:#f2effa;--muted:#a49fb8;--line:rgba(255,255,255,.1)}}
*{box-sizing:border-box}body{margin:0;min-height:100vh;font:15px/1.45 -apple-system,"Helvetica Neue",Helvetica,Arial,sans-serif;color:var(--ink);
background:radial-gradient(60vw 40vh at 10% 0%,rgba(124,92,255,.18),transparent),radial-gradient(50vw 40vh at 100% 100%,rgba(70,200,255,.16),transparent),var(--bg)}
main{max-width:760px;margin:0 auto;padding:28px 16px 60px}h1{font-size:26px;margin:0 0 4px;letter-spacing:-.01em}.sub{color:var(--muted);margin:0 0 18px}
.card{background:var(--card);border:1px solid var(--line);border-radius:20px;padding:16px;backdrop-filter:blur(20px);-webkit-backdrop-filter:blur(20px);box-shadow:0 10px 30px -18px rgba(20,16,40,.35)}
button{font:inherit;border:0;border-radius:999px;padding:8px 14px;cursor:pointer;background:var(--line);color:var(--ink)}button.go{background:linear-gradient(135deg,var(--a),var(--b));color:#fff}
input,select{font:inherit;color:var(--ink);background:transparent;border:1px solid var(--line);border-radius:10px;padding:8px 10px;min-width:0;width:100%}
table{width:100%;border-collapse:collapse}th,td{text-align:left;padding:8px 6px;border-bottom:1px solid var(--line)}th{font-size:12px;text-transform:uppercase;letter-spacing:.04em;color:var(--muted)}
td.n,th.n{text-align:right}.row{display:flex;gap:8px;flex-wrap:wrap;align-items:center}.grow{flex:1}.muted{color:var(--muted)}
.x{background:none;color:var(--muted);padding:2px 8px}.bar{height:8px;border-radius:99px;background:var(--line);overflow:hidden;margin:10px 0 14px}.bar i{display:block;height:100%;background:linear-gradient(90deg,var(--a),var(--b));transition:width .4s}
li{list-style:none;display:flex;gap:10px;align-items:center;padding:8px 4px;border-bottom:1px solid var(--line)}ul{padding:0;margin:0}li.done span{text-decoration:line-through;color:var(--muted)}
.flash{min-height:220px;display:grid;place-items:center;text-align:center;font-size:22px;cursor:pointer;perspective:900px}.big{font-size:56px;font-weight:700;letter-spacing:-.02em}
"""

APP_JS = r"""
const D = JSON.parse(document.getElementById('luma-data').textContent);
const KEY = 'luma-app-' + D.title;
const load = (fallback) => { try { return JSON.parse(localStorage.getItem(KEY)) ?? fallback; } catch (_) { return fallback; } };
const save = (value) => { try { localStorage.setItem(KEY, JSON.stringify(value)); } catch (_) { /* preview: nothing kept */ } };
const h = (tag, attrs = {}, ...kids) => { const e = document.createElement(tag); for (const [k, v] of Object.entries(attrs)) { if (k.startsWith('on')) e.addEventListener(k.slice(2), v); else if (v !== false && v != null) e.setAttribute(k, v); } kids.flat().forEach((k) => e.append(k instanceof Node ? k : document.createTextNode(String(k)))); return e; };
const root = document.getElementById('app');
function tracker() {
  let rows = load([]);
  const draw = () => {
    const form = h('div', { class: 'row' }, D.columns.map((c) => h('input', { placeholder: c.name, type: c.type === 'number' ? 'number' : c.type === 'date' ? 'date' : 'text', step: 'any', 'data-c': c.name, class: 'grow' })),
      h('button', { class: 'go', onclick: () => { const r = {}; form.querySelectorAll('input').forEach((i) => r[i.dataset.c] = i.value); if (Object.values(r).some((v) => v)) { rows.push(r); save(rows); draw(); } } }, 'Add'));
    const nums = D.columns.filter((c) => c.type === 'number');
    const table = h('table', {}, h('tr', {}, D.columns.map((c) => h('th', { class: c.type === 'number' ? 'n' : '' }, c.name)), h('th')),
      rows.map((r, i) => h('tr', {}, D.columns.map((c) => h('td', { class: c.type === 'number' ? 'n' : '' }, r[c.name] || '')), h('td', {}, h('button', { class: 'x', title: 'Remove', onclick: () => { rows.splice(i, 1); save(rows); draw(); } }, '×')))),
      nums.length && rows.length ? h('tr', {}, D.columns.map((c, k) => h('td', { class: c.type === 'number' ? 'n' : 'muted' }, c.type === 'number' ? rows.reduce((s, r) => s + (parseFloat(r[c.name]) || 0), 0).toLocaleString(undefined, { maximumFractionDigits: 2 }) : k === 0 ? 'Total' : '')), h('td')) : '');
    const csv = h('button', { onclick: () => { const lines = [D.columns.map((c) => c.name).join(',')].concat(rows.map((r) => D.columns.map((c) => JSON.stringify(r[c.name] || '')).join(','))); const a = h('a', { href: 'data:text/csv;charset=utf-8,' + encodeURIComponent(lines.join('\n')), download: D.title + '.csv' }); document.body.append(a); a.click(); a.remove(); } }, 'Export CSV');
    root.replaceChildren(h('div', { class: 'card' }, form), h('div', { class: 'card', style: 'margin-top:14px;overflow:auto' }, rows.length ? table : h('p', { class: 'muted' }, 'Nothing yet. Add the first one above.')), h('div', { class: 'row', style: 'margin-top:12px' }, csv));
  };
  draw();
}
function checklist() {
  let items = load(D.items.map((t) => ({ t, done: false })));
  const draw = () => {
    const done = items.filter((i) => i.done).length;
    const input = h('input', { placeholder: 'Add something', class: 'grow', onkeydown: (e) => { if (e.key === 'Enter' && input.value.trim()) { items.push({ t: input.value.trim(), done: false }); save(items); draw(); } } });
    root.replaceChildren(h('div', { class: 'card' }, h('div', { class: 'muted' }, `${done} of ${items.length} done`), h('div', { class: 'bar' }, h('i', { style: `width:${items.length ? 100 * done / items.length : 0}%` })),
      h('ul', {}, items.map((it, i) => h('li', { class: it.done ? 'done' : '' }, h('input', { type: 'checkbox', checked: it.done || false, style: 'width:auto', onchange: () => { it.done = !it.done; save(items); draw(); } }), h('span', { class: 'grow' }, it.t), h('button', { class: 'x', onclick: () => { items.splice(i, 1); save(items); draw(); } }, '×')))),
      h('div', { class: 'row', style: 'margin-top:12px' }, input)));
  };
  draw();
}
function flashcards() {
  let order = D.cards.map((_, i) => i), at = 0, flipped = false, known = load([]);
  const draw = () => {
    const c = D.cards[order[at]];
    root.replaceChildren(h('div', { class: 'muted' }, `Card ${at + 1} of ${order.length} · ${known.length} known`),
      h('div', { class: 'card flash', style: 'margin-top:10px', onclick: () => { flipped = !flipped; draw(); } }, h('div', {}, flipped ? (c.back || '(no answer)') : c.front, h('div', { class: 'muted', style: 'font-size:13px;margin-top:10px' }, flipped ? 'Tap for the question' : 'Tap to see the answer'))),
      h('div', { class: 'row', style: 'margin-top:12px' }, h('button', { onclick: () => { at = (at - 1 + order.length) % order.length; flipped = false; draw(); } }, 'Back'),
        h('button', { onclick: () => { order.sort(() => Math.random() - 0.5); at = 0; flipped = false; draw(); } }, 'Shuffle'), h('span', { class: 'grow' }),
        h('button', { onclick: () => { if (!known.includes(order[at])) known.push(order[at]); save(known); at = (at + 1) % order.length; flipped = false; draw(); } }, 'Got it'),
        h('button', { class: 'go', onclick: () => { at = (at + 1) % order.length; flipped = false; draw(); } }, 'Next')));
  };
  draw();
}
function countdown() {
  const target = new Date(D.date);
  const tick = () => { let s = Math.max(0, Math.round((target - new Date()) / 1000)); const d = Math.floor(s / 86400); s %= 86400; const hh = Math.floor(s / 3600); s %= 3600; const m = Math.floor(s / 60);
    root.replaceChildren(h('div', { class: 'card', style: 'text-align:center;padding:36px 16px' }, h('div', { class: 'big' }, `${d}d ${hh}h ${m}m ${s}s`), h('div', { class: 'muted' }, 'until ' + target.toLocaleString()))); };
  tick(); setInterval(tick, 1000);
}
({ tracker, checklist, flashcards, countdown })[D.kind]();
"""


def app_html(spec):
    data = json.dumps(spec, ensure_ascii=False).replace("</", "<\\/")
    subtitle = {"tracker": "Add a row any time. Totals add themselves up.", "checklist": "Tick things off as you go.",
                "flashcards": "Tap a card to flip it.", "countdown": ""}[spec["kind"]]
    return ("<!doctype html><html lang=en><meta charset=utf-8><meta name=viewport content='width=device-width,initial-scale=1'>"
            f"<title>{html.escape(spec['title'])}</title><style>{APP_CSS}</style><main><h1>{html.escape(spec['title'])}</h1>"
            f"<p class=sub>{html.escape(subtitle)} Made by Luma.</p><div id=app></div></main>"
            f"<script type=application/json id=luma-data>{data}</script><script>{APP_JS}</script></html>")


def build_app(spec, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(app_html(spec), encoding="utf-8")
    return path
