"""The work platform's foundation: project folders nothing can escape, code that runs fenced in,
a search over the project's own files, and brains picked by how hard the job is, with usage counted."""
import json
import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from luma.memory.store import Store
from luma.work import knowledge, sandbox
from luma.work.projects import Projects


class Base(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.store = Store(root / "s.db", root / "k")
        self.projects = Projects(self.store, root / "projects")

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()


class ProjectTests(Base):
    def test_a_project_is_a_folder_with_memory_and_history(self):
        p = self.projects.create("Taco truck launch", "Research a taco truck in Atlanta")
        self.projects.write(p["id"], "research/notes.md", "# Notes\nTacos sell.")
        self.assertEqual([f["path"] for f in self.projects.files(p["id"])], ["research/notes.md"])
        self.assertEqual(self.projects.read(p["id"], "research/notes.md"), "# Notes\nTacos sell.")
        p = self.projects.note(p["id"], "Brand color is teal", "decision")
        self.assertEqual(p["notes"][0]["text"], "Brand color is teal")
        p = self.projects.forget(p["id"], p["notes"][0]["id"])
        self.assertEqual(p["notes"], [])
        self.projects.record_run(p["id"], {"id": "r1", "state": "done"})
        self.assertEqual(self.projects.get(p["id"])["runs"][0]["state"], "done")

    def test_nothing_escapes_the_folder(self):
        p = self.projects.create("x")
        for bad in ["../evil.txt", "/etc/passwd", "a/../../b", "C:/x", ".luma/index.json", ""]:
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                self.projects.write(p["id"], bad, "nope")
        outside = Path(self.temp.name) / "secret.txt"
        outside.write_text("secret")
        os.symlink(outside, self.projects.folder(p["id"]) / "link.txt")
        with self.assertRaises(ValueError):
            self.projects.read(p["id"], "link.txt")

    def test_deleting_needs_a_yes_and_export_packages_everything(self):
        p = self.projects.create("Pitch")
        self.projects.write(p["id"], "deck/notes.txt", "hi")
        self.projects.write(p["id"], "site/index.html", "<h1>hi</h1>")
        with self.assertRaises(PermissionError):
            self.projects.remove(p["id"], "deck/notes.txt")
        with self.assertRaises(PermissionError):
            self.projects.delete(p["id"])
        bundle = self.projects.export(p["id"])
        self.assertEqual(sorted(zipfile.ZipFile(bundle).namelist()), ["deck/notes.txt", "site/index.html"])
        self.projects.delete(p["id"], confirmed=True)
        with self.assertRaises(ValueError):
            self.projects.get(p["id"])


@unittest.skipIf(os.name != "posix", "posix limits")
class SandboxTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.cwd = Path(self.temp.name) / "proj"

    def tearDown(self):
        self.temp.cleanup()

    def test_runs_python_in_the_project_with_a_clean_environment(self):
        self.cwd.mkdir()
        (self.cwd / "hello.py").write_text("import os, json\nprint(json.dumps({'cwd': os.getcwd(), 'keys': sorted(k for k in os.environ if 'KEY' in k or 'TOKEN' in k)}))\nopen('out.txt','w').write('made')\n")
        with mock.patch.dict(os.environ, {"ANTHROPIC_API_KEY": "sk-ant-secret", "GITHUB_TOKEN": "x"}):
            result = sandbox.run("python hello.py", self.cwd)
        self.assertTrue(result["ok"], result)
        said = json.loads(result["stdout"])
        self.assertEqual(Path(said["cwd"]).resolve(), self.cwd.resolve())
        self.assertEqual(said["keys"], [])  # no secrets leak in
        self.assertEqual((self.cwd / "out.txt").read_text(), "made")

    def test_stops_runaway_code_and_reports_failures(self):
        result = sandbox.run(["python", "-c", "while True: pass"], self.cwd, timeout=2)
        self.assertTrue(result["timed_out"])
        self.assertFalse(result["ok"])
        failed = sandbox.run(["python", "-c", "raise SystemExit('broke')"], self.cwd)
        self.assertFalse(failed["ok"])
        self.assertIn("broke", failed["stderr"])

    def test_only_python_and_node_and_no_installs(self):
        for bad in ["rm -rf /", "bash -c ls", "curl https://example.com", "python -m pip install requests", "python /etc/passwd"]:
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                sandbox.run(bad, self.cwd)

    def test_apple_sandbox_rules_fence_in_the_folder(self):
        rules = sandbox.profile("/Users/me/Luma/projects/p1", ["/opt/python"])
        self.assertIn("(deny network*)", rules)
        self.assertLess(rules.index("(deny file-write*)"), rules.index('(allow file-write* (subpath "/Users/me/Luma/projects/p1")'))
        self.assertLess(rules.index("(deny file-read*"), rules.index('(allow file-read* (subpath "/Users/me/Luma/projects/p1")'))
        self.assertIn('(subpath "/opt/python")', rules)


class KnowledgeTests(unittest.TestCase):
    def test_finds_the_passage_and_cites_the_file(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            (folder / "market.md").write_text("Atlanta food trucks\n\nThe Atlanta food truck market grew 12 percent last year, led by tacos and wings.\n\nPermits cost about $500.")
            (folder / "team.txt").write_text("Our team has two chefs and one designer.")
            from docx import Document
            doc = Document()
            doc.add_paragraph("Competitor analysis: Tacos Chingon charges nine dollars per plate.")
            doc.save(str(folder / "competitors.docx"))
            index = knowledge.Index(folder).refresh()
            hits = index.search("how much did the food truck market grow")
            self.assertEqual(hits[0]["file"], "market.md")
            self.assertIn("12 percent", hits[0]["text"])
            self.assertEqual(index.search("competitor plate price")[0]["file"], "competitors.docx")
            self.assertEqual(knowledge.Index(folder).search("chefs")[0]["file"], "team.txt")  # saved index loads


class DeckChartTests(unittest.TestCase):
    def test_charts_diagrams_and_tables_are_real_and_checked(self):
        from pptx import Presentation
        from luma.skills import decks
        spec = {"title": "Taco Truck", "theme": "ember", "slides": [
            {"layout": "chart", "title": "Revenue doubles every year", "chart_type": "column", "categories": ["2026", "2027"],
             "series": [{"name": "Revenue", "values": [120000, 260000]}]},
            {"layout": "chart", "title": "Where sales come from", "chart_type": "pie", "categories": ["Lunch", "Dinner"], "series": [{"name": "Share", "values": [60, 40]}]},
            {"layout": "diagram", "title": "How it works", "steps": ["Order", "Cook", "Pick up"]},
            {"layout": "table", "title": "Competitors", "rows": [["Name", "Price"], ["Tacos Chingon", "$9"]]}]}
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder) / "d.pptx"
            self.assertEqual(decks.build_deck(spec, out), 4)
            slides = Presentation(str(out)).slides
            self.assertEqual(sum(1 for s in slides for sh in s.shapes if sh.has_chart), 2)
            self.assertEqual(sum(1 for s in slides for sh in s.shapes if sh.has_table), 1)
        bad = {"title": "x", "slides": [{"layout": "chart", "title": "t", "chart_type": "line", "categories": ["a", "b"], "series": [{"name": "s", "values": [1]}]}]}
        with self.assertRaises(ValueError):
            decks.validate_outline(bad)  # one value per category, or it doesn't get drawn


class RoutingTests(Base):
    def tearDown(self):
        from luma.llm import brain
        brain._store = None
        super().tearDown()

    def test_counting_never_breaks_a_call(self):
        from luma.llm import usage
        self.store.close()
        usage.record(self.store, "x", 1, 1, 0.1)  # a closed store: nothing raised
        self.store = Store(Path(self.temp.name) / "s.db", Path(self.temp.name) / "k")

    def test_light_jobs_stay_local_and_usage_is_counted(self):
        from luma.llm import brain, inference, usage
        brain.init(self.store)
        self.store.set_setting("brain_cloud", True)
        self.store.set_setting("brain_cloud_key", "sk-ant-x")
        with mock.patch.object(brain, "health", return_value={"ok": True}):
            self.assertFalse(brain.use_cloud("light"))
            self.assertTrue(brain.use_cloud("heavy"))
            self.store.set_setting("route_by_task", False)
            self.assertTrue(brain.use_cloud("light"))
        self.store.set_setting("brain_cloud", False)
        with mock.patch.object(inference, "plan", return_value={"type": "reply", "text": "hey there"}):
            brain.plan([{"role": "user", "content": "hi"}], [], {}, "friend", tier="light")
        totals = usage.summary(self.store)["models"]
        (label, row), = totals.items()
        self.assertIn("local", label)
        self.assertEqual(row["calls"], 1)
        self.assertGreater(row["tokens_in"], 0)
        self.assertIsNone(row["cost"])  # no price set: no guess
        self.store.set_setting("model_prices", {label: [1.0, 2.0]})
        self.assertIsNotNone(usage.summary(self.store)["models"][label]["cost"])


if __name__ == "__main__":
    unittest.main()
