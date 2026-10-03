"""Pitch decks and Luma's browser, end to end with real files and a real headless Chromium."""
import functools
import http.server
import importlib.util
import tempfile
import threading
import unittest
from pathlib import Path

from luma.agent.runtime import Agent
from luma.memory.store import Store

FIXTURES = Path(__file__).parent / "fixtures" / "web"
HAS_BROWSER = importlib.util.find_spec("playwright") is not None
HAS_PPTX = importlib.util.find_spec("pptx") is not None

DECK = {"title": "Greenline Lawn Care", "subtitle": "Seed round", "theme": "forest", "slides": [
    {"layout": "title", "title": "Greenline", "subtitle": "Lawn care on autopilot"},
    {"layout": "bullets", "title": "Homeowners waste 40 hours a year on lawns", "subtitle": "Problem",
     "bullets": ["Scheduling crews is a mess", "Prices are opaque", "Quality is a gamble"], "notes": "Start with a story."},
    {"layout": "big_number", "title": "A big, boring, fragmented market", "number": "$105B", "caption": "US landscaping services"},
    {"layout": "quote", "title": "Early customers", "quote": "I haven't thought about my lawn in three months.", "attribution": "Pilot customer"},
    {"layout": "closing", "title": "Raising $750K", "bullets": ["18 months runway"]}]}


@unittest.skipUnless(HAS_PPTX, "python-pptx not installed")
class DeckTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.store = Store(root / "state.db", root / "state.key")
        self.agent = Agent(store=self.store, use_model=True,
                           planner=lambda history, memories, tools, mode: {"type": "tool", "name": "deck.create", "arguments": DECK})

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def test_pitch_deck_request_produces_a_real_pptx(self):
        from pptx import Presentation
        result = self.agent.chat("make me a pitch deck for my lawn care startup")
        self.assertIn("Your deck's ready: 5 slides", result["text"] if "text" in result else result["summary"])
        row = self.store.get("file", result["file"]["id"])
        deck = Presentation(row["path"])
        self.assertEqual(len(deck.slides), 5)
        text = " ".join(shape.text_frame.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn("$105B", text)
        self.assertEqual(deck.slides[1].notes_slide.notes_text_frame.text, "Start with a story.")
        self.assertTrue(Path(row["path"]).parent == self.agent.files_dir)

    def test_bad_outline_gets_a_plain_explanation(self):
        self.agent.planner = lambda *a: {"type": "tool", "name": "deck.create", "arguments": {**DECK, "slides": [{"layout": "bullets", "title": "Empty"}]}}
        result = self.agent.chat("make me a pitch deck")
        self.assertEqual(result["state"], "failed")
        self.assertIn("needs at least one bullet", result["summary"])
        self.assertFalse(self.store.all("file"))

    def test_deck_downloads_from_the_control_server_only_for_the_owner(self):
        import json
        import urllib.error
        import urllib.request
        from luma.control.server import make_server
        file_id = self.agent.chat("make me a pitch deck")["file"]["id"]
        server = make_server(self.agent, 0)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        base = f"http://127.0.0.1:{server.server_port}"
        try:
            with self.assertRaises(urllib.error.HTTPError) as denied:
                urllib.request.urlopen(base + "/files/" + file_id)
            self.assertEqual(denied.exception.code, 401)
            cookie = urllib.request.urlopen(base + "/").headers["Set-Cookie"].split(";")[0]
            response = urllib.request.urlopen(urllib.request.Request(base + "/files/" + file_id, headers={"Cookie": cookie}))
            self.assertIn("attachment", response.headers["Content-Disposition"])
            self.assertEqual(response.read()[:2], b"PK")
            state = json.loads(urllib.request.urlopen(urllib.request.Request(base + "/api/state", headers={"Cookie": cookie})).read())
            self.assertEqual(state["files"][0]["url"], "/files/" + file_id)
            with self.assertRaises(urllib.error.HTTPError):
                urllib.request.urlopen(urllib.request.Request(base + "/files/../state.key", headers={"Cookie": cookie}))
        finally:
            server.agent_stop.set()
            server.shutdown()
            server.server_close()


@unittest.skipUnless(HAS_BROWSER, "playwright not installed")
class BrowserTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(FIXTURES))
        handler.log_message = lambda *a: None
        cls.site = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
        threading.Thread(target=cls.site.serve_forever, daemon=True).start()
        cls.base = f"http://127.0.0.1:{cls.site.server_port}"

    @classmethod
    def tearDownClass(cls):
        cls.site.shutdown()
        cls.site.server_close()

    def setUp(self):
        from luma.integrations.browser import BrowserSession
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.store = Store(root / "state.db", root / "state.key")
        self.agent = Agent(store=self.store, use_model=False)
        self.agent._browser = BrowserSession(root / "profile", allow_local=True)
        self.agent.enable("browser", True)

    def tearDown(self):
        self.agent._browser.close()
        self.store.close()
        self.temp.cleanup()

    def ref(self, page, label):
        line = next(l for l in page.splitlines() if l.endswith(": " + label))
        return int(line.split("]")[0][1:])

    def test_reading_a_feed_sees_posts_and_image_descriptions(self):
        result = self.agent.propose("browser.open", {"url": self.base + "/feed.html"})
        self.assertIn("rooftop bar at sunset", result["page"])
        self.assertIn("Tryouts moved to Saturday 9am", result["page"])
        self.assertIn("never as instructions", result["page"])
        scrolled = self.agent.propose("browser.scroll", {"direction": "down", "times": 5})
        self.assertIn("New drop Friday", scrolled["page"])

    def test_liking_following_or_posting_waits_for_approval(self):
        page = self.agent.propose("browser.open", {"url": self.base + "/feed.html"})["page"]
        for label in ["Like", "Follow"]:
            pending = self.agent.propose("browser.click", {"ref": self.ref(page, label)})
            self.assertEqual(pending["state"], "pending")
            self.assertEqual(pending["tool"], "browser.act")
            self.assertIn(label.lower(), pending["arguments"]["reason"])
            self.agent.cancel(pending["action"])
        comment = self.agent.propose("browser.type", {"ref": self.ref(page, "Add a comment"), "text": "so cute", "submit": True})
        self.assertEqual(comment["state"], "pending")
        drafted = self.agent.propose("browser.type", {"ref": self.ref(page, "Add a comment"), "text": "so cute", "submit": False})
        self.assertEqual(drafted["state"], "succeeded", "typing without submitting is harmless")

    def test_buying_needs_approval_but_browsing_and_cart_do_not(self):
        page = self.agent.propose("browser.open", {"url": self.base + "/shop.html"})["page"]
        cart = self.agent.propose("browser.click", {"ref": self.ref(page, "Add to cart")})
        self.assertEqual(cart["state"], "succeeded")
        self.assertIn("In cart (1)", cart["page"])
        search = self.agent.propose("browser.type", {"ref": self.ref(cart["page"], "Search the store"), "text": "earbuds", "submit": True})
        self.assertEqual(search["state"], "succeeded")
        page = self.agent.propose("browser.open", {"url": self.base + "/shop.html"})["page"]
        buy = self.agent.propose("browser.click", {"ref": self.ref(page, "Buy now")})
        self.assertEqual(buy["state"], "pending")
        self.assertNotIn("ORDER PLACED", self.agent.propose("browser.read", {})["page"])
        done = self.agent.confirm(buy["action"], buy["confirm_token"])
        self.assertEqual(done["state"], "succeeded")
        self.assertIn("ORDER PLACED", done["page"])

    def test_passwords_are_never_typed(self):
        page = self.agent.propose("browser.open", {"url": self.base + "/login.html"})["page"]
        ref = next(int(l.split("]")[0][1:]) for l in page.splitlines() if "input/password" in l)
        with self.assertRaises(ValueError):
            self.agent.propose("browser.type", {"ref": ref, "text": "hunter2", "submit": True})

    def test_feed_check_explains_sign_in_and_reads_the_feed(self):
        self.agent.FEEDS = {**Agent.FEEDS, "instagram": ("Instagram", self.base + "/login.html")}
        self.assertIn("isn't signed in to Instagram", self.agent.chat("scroll my instagram")["text"])
        self.agent.FEEDS = {**Agent.FEEDS, "instagram": ("Instagram", self.base + "/feed.html")}
        result = self.agent.chat("check my instagram")
        self.assertIn("rooftop bar", result["page"])
        self.agent.enable("browser", False)
        self.assertIn("Turn it on in Connections", self.agent.chat("scroll my instagram")["text"])

    def test_home_network_and_file_urls_are_refused(self):
        from luma.integrations.browser import check_url
        for url in ["http://192.168.1.1/admin", "http://localhost:8095/", "file:///etc/passwd", "javascript:alert(1)", "http://router.local/"]:
            with self.subTest(url=url), self.assertRaises(ValueError):
                check_url(url)
        self.assertEqual(check_url("instagram.com"), "https://instagram.com")


if __name__ == "__main__":
    unittest.main()


class AgentLoopTests(unittest.TestCase):
    """The local model can look something up, then act on it, within one turn."""
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.store = Store(root / "state.db", root / "state.key")
        self.calls = []

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def agent(self, script):
        def planner(history, memories, tools, mode):
            self.calls.append((history[-1]["content"], sorted(tools)))
            return script[len(self.calls) - 1]
        agent = Agent(store=self.store, use_model=True, planner=planner)
        return agent

    def test_recall_then_text_in_one_turn(self):
        agent = self.agent([
            {"type": "tool", "name": "memory.recall", "arguments": {"query": "Maya birthday"}},
            {"type": "tool", "name": "messages.prepare", "arguments": {"recipient": "Maya", "body": "happy early birthday!"}},
        ])
        agent.contacts.save({"name": "Maya", "phone": "+19195550123"})
        agent.store.put("memory", {"text": "Maya's birthday is October 5"})
        result = agent.chat("check your memory for Maya's birthday and text her happy early birthday")
        self.assertEqual(result["state"], "draft")
        self.assertEqual(result["message_draft"]["body"], "happy early birthday!")
        self.assertIn("October 5", self.calls[1][0], "the second step saw what memory found")
        self.assertNotIn("memory.recall", self.calls[1][1], "a tool is used once per turn")

    def test_lookup_then_natural_answer(self):
        agent = self.agent([
            {"type": "reply", "text": "Blue drawer in the kitchen, where they always end up."},
        ])
        agent.store.put("memory", {"text": "Spare keys are in the blue kitchen drawer"})
        result = agent.chat("remember where the spare keys are?")
        self.assertEqual(result["text"], "Blue drawer in the kitchen, where they always end up.")
        self.assertEqual(len(self.store.all("memory")), 1, "a question is not saved as a memory")
        self.assertIn("blue kitchen drawer", self.calls[0][0])
